import io, json, zipfile
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from sejong_rag.extract_attachments import _safe_zip_members, extract_docx, extract_one, extract_pptx, extract_xlsx

def test_extract_xlsx_preserves_sheet_and_row():
    stream=io.BytesIO(); book=Workbook(); sheet=book.active; sheet.title="졸업요건"; sheet.append(["학과","필요학점"]); sheet.append(["소프트웨어학과",130]); book.save(stream); stream.seek(0)
    records,warnings=extract_xlsx(stream); assert warnings==[]; assert records[1]["locator"]=={"sheet":"졸업요건","row":2}; assert "130" in records[1]["text"]

def test_extract_docx_reads_paragraph_and_table():
    stream=io.BytesIO(); document=Document(); document.add_paragraph("장학금 신청 안내"); table=document.add_table(rows=1,cols=2); table.cell(0,0).text="마감"; table.cell(0,1).text="8월 31일"; document.save(stream); stream.seek(0)
    records,_=extract_docx(stream); assert any("장학금" in x["text"] for x in records); assert any(x["locator"].get("table")==1 for x in records)

def test_extract_pptx_preserves_slide():
    stream=io.BytesIO(); presentation=Presentation(); slide=presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text="수강신청"; presentation.save(stream); stream.seek(0)
    records,_=extract_pptx(stream); assert records[0]["locator"]=={"slide":1}

def test_zip_rejects_path_traversal(tmp_path: Path):
    path=tmp_path/"bad.zip"
    with zipfile.ZipFile(path,"w") as archive: archive.writestr("../escape.txt","bad")
    with zipfile.ZipFile(path) as archive:
        try: _safe_zip_members(archive)
        except ValueError as exc: assert "unsafe ZIP path" in str(exc)
        else: raise AssertionError("unsafe ZIP path was accepted")

def test_legacy_doc_is_failed_explicitly(tmp_path: Path):
    raw=tmp_path/"sample.doc"; raw.write_bytes(b"legacy"); metadata={"document_id":"a-1","resolved_raw_path":str(raw),"file_type":"doc","filename":"sample.doc"}
    report=extract_one(metadata,tmp_path/"out"); assert report["status"]=="failed"; assert "not supported safely" in report["error"]
    assert json.loads((tmp_path/"out/reports/a-1.json").read_text())["record_count"]==0
