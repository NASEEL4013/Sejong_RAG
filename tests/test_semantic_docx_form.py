from pathlib import Path
import pytest
from docx import Document
from sejong_rag.extract_attachments import extract_docx
from sejong_rag.semantic_docx_form import topic_plans,FormLayoutReview
from sejong_rag.semantic_rechunk import split_table_form,digest

class Helper:
    def count(self,text):return len(text)//2+2

def fixture(tmp_path,limit='500',period='2025년'):
    d=Document();d.add_paragraph('장학생 신청서')
    t=d.add_table(rows=7,cols=4)
    for i,text in enumerate(['성명','','학번','']):t.cell(0,i).text=text
    t.cell(1,0).merge(t.cell(1,3)).text='첨부서류 확인표'
    t.cell(2,0).merge(t.cell(2,1)).text='성적증명서'
    t.cell(2,2).merge(t.cell(2,3)).text='성적증명서'
    t.cell(3,0).merge(t.cell(3,3)).text=f'납부확인서 발급기간: {period}'
    t.cell(4,0).merge(t.cell(4,3)).text='1. 신청 동기를 서술해주세요.'
    t.cell(5,0).merge(t.cell(5,3)).text=f'- 띄어쓰기 제외 {limit}자 이내 -'
    t.cell(6,0).merge(t.cell(6,3)).text='추가 제출 안내'
    p=tmp_path/'form.docx';d.save(p)
    raw,_=extract_docx(p)
    rs=[dict(r,record_id=f'r{i}',document_id='d',document_title='장학생 선발',filename='신청서.docx',file_type='docx') for i,r in enumerate(raw)]
    return p,rs

def test_original_cell_identity_not_equal_text_controls_deduplication(tmp_path):
    p,rs=fixture(tmp_path)
    plans=topic_plans(p,rs)
    attachments=next(x for x in plans if x['title']=='첨부서류 확인표')
    assert '납부확인서 발급기간: 2025년' in attachments['body']
    assert attachments['body'].count('성적증명서')==2 # two distinct merged cells
    assert attachments['body'].count('첨부서류 확인표')==2 # topic label and source heading
    q=next(x for x in plans if x['title'].startswith('1.'))
    assert q['body'].count('500자 이내')==1
    assert '추가 제출 안내' not in q['body']
    chunks=split_table_form(rs,{r['record_id']:r for r in rs},Helper(),p,digest(p))
    assert {s['record_id'] for c in chunks for s in c['context_source_spans']}=={r['record_id'] for r in rs}
    assert all(c['table_geometry']['source_sha256']==digest(p) for c in chunks)

def test_changed_limits_and_periods_come_from_source(tmp_path):
    p,rs=fixture(tmp_path,'700','2026년')
    text='\n'.join(x['body'] for x in topic_plans(p,rs))
    assert '700자' in text and '2026년' in text and '500자' not in text

def test_normalized_mismatch_and_missing_records_are_rejected(tmp_path):
    p,rs=fixture(tmp_path);rs[-1]['text']='변조된 내용'
    with pytest.raises(FormLayoutReview,match='row_source_mismatch'):topic_plans(p,rs)
    with pytest.raises(FormLayoutReview,match='row_source_mismatch'):topic_plans(p,rs[:-1])

def test_limit_must_remain_attached_to_question(tmp_path):
    p,rs=fixture(tmp_path);d=Document(p)
    d.tables[0].cell(5,0).text='별도 안내'
    d.save(p);raw,_=extract_docx(p)
    for r,item in zip(rs,raw):r['text']=item['text']
    with pytest.raises(FormLayoutReview,match='question_limit_missing'):topic_plans(p,rs)


def test_unknown_full_width_text_is_not_assumed_to_be_a_heading(tmp_path):
    p,rs=fixture(tmp_path);d=Document(p)
    d.tables[0].cell(5,0).text='모호한 새로운 내용'
    d.save(p);raw,_=extract_docx(p)
    for r,item in zip(rs,raw):r['text']=item['text']
    with pytest.raises(FormLayoutReview,match='full_width_row_requires_review'):topic_plans(p,rs)
