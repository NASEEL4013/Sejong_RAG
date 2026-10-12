"""Export original 2026-10-11 benchmark as a browsable HTML and Excel workbook."""
from __future__ import annotations
import hashlib
import html
import json
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'data/processed/project-benchmark-v1-resumed/report.json'
OUT=ROOT/'reports'
STEM='benchmark-review-2026-10-11'
STATUS={'answered':'답변 제공','insufficient_evidence':'근거 부족으로 보류','needs_clarification':'조건 확인 요청'}
KIND={'answerable':'답변 가능','unanswerable':'근거 없음','clarification':'조건 부족','privacy':'개인정보·개인 판정','adversarial':'허위 답변 요구'}
METRICS=[('상위 5개 통합 정답 근거','19/24 · 79.2%','Top5 전체를 합쳐 질문의 모든 항목을 지원하는지 판단'),
 ('완전한 답변 정확도','19/24 · 79.2%','보류·일부 조건 누락도 실패에 포함'),
 ('실제 답변 제공','20/24 · 83.3%','답변 가능한 24문항 중 answered 상태'),
 ('답변한 질문 중 정확도','19/20 · 95%','답변 제공 문항만을 분모로 계산'),
 ('주장별 인용 근거 지원','46/47 · 97.9%','claims 배열의 각 주장을 인용 본문과 대조'),
 ('보류·조건 확인 등 행동 평가','8/8 · 100%','근거 없음·조건 부족·개인정보·허위 요구 대응')]


def esc(value):return html.escape(str(value if value is not None else ''))
def href(url):
    if urlparse(str(url or '')).scheme not in ('http','https'):return ''
    return esc(url)
def link(url,label):
    return f'<a href="{href(url)}" target="_blank" rel="noopener noreferrer">{esc(label)} ↗</a>' if href(url) else ''
def body_preview(text):return ' '.join(text.split())[:170]
def location(bundle):
    result=[]
    for loc in bundle.get('locations',bundle.get('locators',[])):
        if 'page' in loc:result.append('페이지 '+str(loc['page']))
        elif loc.get('section')=='notice_body':result.append('공지 본문')
        else:result.append(json.dumps(loc,ensure_ascii=False))
    return ' / '.join(dict.fromkeys(result)) or '원문 구간 기록 참조'


def load():
    raw=REPORT.read_bytes(); report=json.loads(raw); digest=hashlib.sha256(raw).hexdigest()
    reviews={}
    for name in ['project_benchmark_review_a_v1.json','project_benchmark_review_b_v1.json']:
        review=json.loads((ROOT/'eval'/name).read_text())
        assert review['report_sha256']==digest
        for item in review['results']:
            assert item['id'] not in reviews
            reviews[item['id']]=item
    amendments=json.loads((ROOT/'eval/project_benchmark_amendments_v1.json').read_text())
    assert amendments['report_sha256']==digest
    for amendment in amendments['results']:
        review=reviews[amendment['id']]
        for key in ('answer_correct','failure_types','reason'):
            review[key]=amendment[key]
        for index,reason in amendment.get('claim_reason_overrides',{}).items():
            review['claim_checks'][int(index)-1]['reason']=reason
        row=next(row for row in report['results'] if row['id']==amendment['id'])
        for key in ('answer_note','rubric'):
            row[key]=amendment[key]
    assert len(report['results'])==32 and len(reviews)==32
    union=json.loads((ROOT/'eval/project_benchmark_retrieval_review_v2.json').read_text())
    assert union['source_report_sha256']==digest
    union_by_id={r['id']:r for r in union['results']}
    for row in report['results']:
        row['retrieval_review']=union_by_id.get(row['id'])
        row['review']=reviews[row['id']]
        assert len(row['current']['top5'])==5
        assert len(row['generation']['result']['claims'])==len(row['review']['claim_checks'])
        row['passed']=row['review']['answer_correct'] if row['kind']=='answerable' else row['review']['behavior_correct']
        row['citation_issue']=any(not c['supported'] for c in row['review']['claim_checks']) or row['review'].get('message_fact_leak',False)
    assert sum(r['passed'] for r in report['results'] if r['kind']=='answerable')==19
    return report,digest


def make_html(report,digest):
    cards=[]
    for row in report['results']:
        ident=row['id']; gen=row['generation']['result']; review=row['review']; sources=row['current']['top5']
        first=row['current']['first_relevant_rank']; cited={i for c in gen['claims'] for i in c['citations']}
        grade=('완전 정답' if row['kind']=='answerable' else '행동 통과') if row['passed'] else '완전 답변 실패'
        tags=[f'<span class="badge {"ok" if row["passed"] else "bad"}">{grade}</span>',
              f'<span class="badge">{esc(KIND[row["kind"]])}</span>',f'<span class="badge">{esc(STATUS[gen["status"]])}</span>']
        if row['citation_issue']:tags.append('<span class="badge warn">인용·근거 문제</span>')
        claims=''
        for n,claim in enumerate(gen['claims'],1):
            checked=review['claim_checks'][n-1]
            labels=' '.join(f'<button class="cite" data-target="{ident}-r{i}">근거 {i} 보기</button>' for i in claim['citations'])
            claims+=f'<li><p>{esc(claim["text"])}</p>{labels}<span class="support {"oktext" if checked["supported"] else "badtext"}">인용 지원 {"확인" if checked["supported"] else "안 됨"}</span><div class="small">{esc(checked["reason"])}</div></li>'
        gold=''.join(f'<p><a href="../{esc(s["path"])}" target="_blank" rel="noopener">저장된 정답 원문 · {esc(s["notice_id"])}</a></p>' + (f'<blockquote>{esc(" / ".join(s.get("quotes",[])))}</blockquote>' if s.get('quotes') else '') for s in row.get('gold_sources',[]))
        retrieval=[]
        for b in sources:
            rank=b['rank']; used=rank in cited
            bundle_tags=f'<span class="badge {"ok" if b["is_evidence"] else ""}">{"단일 묶음 정답 일치" if b["is_evidence"] else "단일 묶음 기준 미충족"}</span>'
            if row['kind']!='answerable':bundle_tags='<span class="badge">검색 정답 채점 대상 아님</span>'
            if used:bundle_tags+='<span class="badge blue">모델이 인용함</span>'
            retrieval.append(f'''<details class="evidence" id="{ident}-r{rank}"><summary><span class="rank">{rank}</span><span class="source-summary"><strong>{esc(b['document_title'])}</strong><span class="small">{esc(b.get('filename') or '공지 본문')} · {esc(location(b))} · 유사도 {b['score']:.4f} · {b['token_count']} E5 토큰</span><span class="badges">{bundle_tags}</span><span class="preview">{esc(body_preview(b['body']))}</span></span></summary><div class="evidence-body"><div class="links">{link(b.get('source_url'),'공식 공지')}{link(b.get('attachment_url'),'첨부파일')}</div><h4>검색된 본문 전체</h4><pre>{esc(b['body'])}</pre><details><summary>제목을 포함한 모델 전달 텍스트</summary><pre>{esc(b['text'])}</pre></details><details><summary>청크·원문 위치 추적 정보</summary><pre>{esc(json.dumps({k:b.get(k) for k in ['anchor_chunk_id','chunk_ids','source_spans','context_expanded']},ensure_ascii=False,indent=2))}</pre></details></div></details>''')
        leak='<p class="alert">보류 message에 인용 없는 사실이 있어. 이 내용은 47개 claim의 인용 검사 분모에 포함되지 않았어.</p>' if review.get('message_fact_leak') else ''
        union=row.get('retrieval_review')
        union_note=('통합 검색 성공' if union['complete'] else '통합 검색 미충족')+' — '+union['reason'] if union else '행동 평가 문항'
        single=f'{first}위' if first else ('해당 없음 — 행동 평가' if row['kind']!='answerable' else '단일 묶음 일치 없음')
        cards.append(f'''<article class="case" id="{ident}" data-failed="{str(not row['passed']).lower()}" data-retrieval-failed="{str(row['kind']=='answerable' and not row['retrieval_review']['complete']).lower()}" data-behavior="{str(row['kind']!='answerable').lower()}" data-citation="{str(row['citation_issue']).lower()}"><header class="case-head"><span class="case-id">{ident}</span><div><h2>{esc(row['query'])}</h2><div class="badges">{''.join(tags)}</div></div></header><div class="twocol"><section><h3>실제 모델 답변</h3><div class="answer">{esc(gen['answer'])}</div>{leak}<details><summary>주장별 인용 검토 ({len(gen['claims'])}개)</summary><ol class="claims">{claims or '<li>claims가 비어 있어. 위 message로 응답했어.</li>'}</ol></details></section><section><h3>정답·기대 행동</h3><p>{esc(row['answer_note'])}</p><h4>판정 기준</h4><p>{esc(row.get('rubric',''))}</p><h4>이렇게 판정한 이유</h4><p>{esc(review['reason'])}</p><details><summary>정답 원문·확인 문구</summary>{gold or '<p>보류·조건 확인 등 행동 평가 문항이야.</p>'}</details></section></div><div class="retrieval-head"><h3>검색 상위 5개</h3><p>{esc(union_note)}</p><span>참고: 단일 묶음 완전 근거 첫 순위: <strong>{single}</strong> · 답변 인용: {', '.join(map(str,sorted(cited))) or '없음'}</span></div>{''.join(retrieval)}<footer>검색 방식 vector_context · {sum(b['token_count'] for b in sources)} / 2,200 E5 토큰 · 답변 생성 {row['generation']['wall_ms']/1000:.2f}초 · API 입력 {gen['usage']['prompt_token_count']:,} / 출력 {gen['usage']['candidates_token_count']:,}토큰</footer></article>''')
    metrics=''.join(f'<div class="metric"><span>{esc(name)}</span><strong>{value}</strong><small>{esc(note)}</small></div>' for name,value,note in METRICS)
    return '''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>세종 RAG · 질문·검색·답변 검토</title><style>
:root{--ink:#172b3c;--muted:#566b7b;--line:#dbe4ea;--blue:#075fa5}*{box-sizing:border-box}body{margin:0;background:#f3f6f9;color:var(--ink);font:15px/1.65 system-ui,-apple-system,'Malgun Gothic',sans-serif}main{max-width:1220px;margin:auto;padding:32px 24px 70px}h1{font-size:30px;margin:4px 0 12px;line-height:1.3}h2{font-size:19px;margin:0 0 12px;line-height:1.6}h3{font-size:16px;margin:0 0 12px}h4{font-size:14px;margin:16px 0 6px}p{margin:8px 0}.subtitle,.small,footer,small{color:var(--muted)}.small{font-size:12px;display:block}a{color:var(--blue);overflow-wrap:anywhere}a:hover{text-decoration:none}.metrics{display:grid;grid-template-columns:repeat(3,1fr);gap:12px;margin:24px 0}.metric{background:white;padding:18px;border:1px solid var(--line);border-radius:12px}.metric span{font-size:13px}.metric strong{display:block;font-size:25px;margin:5px 0}.metric small{font-size:12px}.notice{border-left:4px solid #4d7a9a;padding:12px 18px;background:#e8f0f6;font-size:13px}.toolbar{position:sticky;top:0;z-index:2;background:#f3f6f9f5;display:flex;gap:10px;align-items:center;flex-wrap:wrap;padding:18px 0;border-bottom:1px solid var(--line);margin-bottom:20px}input,select,button{font:inherit;border:1px solid #bacbd7;border-radius:7px;background:white;padding:8px 12px;color:var(--ink)}input{flex:1;min-width:200px}button{cursor:pointer}button:hover{background:#e8f0f6}.case{background:white;border:1px solid var(--line);border-radius:12px;padding:24px;margin:0 0 24px;scroll-margin-top:95px}.case[hidden]{display:none}.case-head{display:flex;gap:16px;padding-bottom:18px;border-bottom:1px solid var(--line)}.case-id{font-weight:700;background:#edf3f8;padding:5px 10px;height:36px;border-radius:6px}.badges{display:flex;gap:6px;flex-wrap:wrap}.badge{font-size:11px;font-weight:600;background:#eef2f5;padding:3px 8px;border-radius:5px;display:inline-block}.ok{background:#e3f4e9;color:#166d3f}.bad{background:#fde8e6;color:#9f251a}.warn{background:#fff1d5;color:#805407}.blue{background:#e2edff;color:#1a4d98}.twocol{display:grid;grid-template-columns:1fr 1fr;gap:24px;padding:22px 0}.answer{white-space:pre-wrap;overflow-wrap:anywhere;background:#f4f8fc;border:1px solid #e0e9f2;border-radius:8px;padding:16px}.claims{padding-left:22px}.claims li{margin:15px 0}.cite{font-size:12px;margin:0 5px 5px 0;padding:3px 7px}.support{font-size:12px}.oktext{color:#166d3f}.badtext{color:#ad2d22}.alert{background:#fff1d5;padding:10px;font-size:13px;border-radius:6px}.retrieval-head{display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap;margin:12px 0}.retrieval-head span{font-size:12px;color:var(--muted)}details>summary{cursor:pointer;color:#244c6a;font-size:13px;padding:9px 0}.evidence{border:1px solid var(--line);border-radius:8px;margin:9px 0}.evidence>summary{display:flex;gap:12px;padding:14px;list-style:none}.rank{background:#eaf1f7;border-radius:6px;font-weight:700;width:30px;height:30px;text-align:center;flex-shrink:0;line-height:30px}.source-summary{display:block;min-width:0}.source-summary strong{font-size:14px;display:block}.preview{display:block;color:#637684;font-size:12px;margin-top:7px}.evidence[open]>summary{border-bottom:1px solid var(--line)}.evidence[open] .preview{display:none}.evidence-body{padding:4px 18px 18px}.links{display:flex;gap:16px;font-size:13px}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f7f9fb;border:1px solid #e6edf2;border-radius:7px;padding:14px;font:13px/1.75 system-ui,sans-serif;max-height:620px;overflow:auto}blockquote{margin:10px 0;padding:10px 15px;border-left:3px solid #c9d9e5;font-size:13px}footer{font-size:11px;padding-top:12px}.provenance{font-size:11px;overflow-wrap:anywhere;margin-top:30px}#count{font-size:13px}.empty{padding:30px;text-align:center}button:focus-visible,summary:focus-visible,input:focus-visible,select:focus-visible{outline:3px solid #237db9;outline-offset:3px}@media(max-width:760px){main{padding:20px 12px}.metrics{grid-template-columns:repeat(2,1fr)}.twocol{grid-template-columns:1fr}.case{padding:16px}.case-head{gap:8px}.metric strong{font-size:20px}.toolbar{position:static}h1{font-size:25px}}@media print{body{background:white}.toolbar{display:none}.case{break-before:page;border:0}.twocol{grid-template-columns:1fr}pre{max-height:none}.metrics{grid-template-columns:repeat(3,1fr)}details{break-inside:avoid}}
</style></head><body><main><div class="subtitle">2026-10-11 원본 평가 · 32문항 / 검색 근거 160개</div><h1>어떤 질문에, 어떤 근거로 답했나</h1><p>질문 → 실제 모델 답변 → 정답 기준 → 검색 상위 5개 순서로 확인할 수 있어.</p><p class="subtitle">모델: gemini-3.5-flash-lite · 검색: vector_context · 이후 분해 검색 실험과는 별도인 원본 평가야.</p><div class="metrics">'''+metrics+'''</div><div class="notice"><strong>읽는 법</strong><br>검색 근거 행을 누르면 본문 전체가 열려. ‘모델이 인용함’은 실제 답변에서 사용한 근거야.<br>‘단일 묶음 기준 미충족’은 해당 묶음 하나가 정답 전체를 충족하지 않는다는 뜻이야. 상단 검색 지표와 필터는 Top5 전체를 합쳐 채점했으며, B07·B10·B15·B17도 성공이야. 유사도 점수는 정답 확률이 아니야.<br>판정은 assistant의 원문 대조야. 2026-10-12 사용자 정정으로 B10을 정답 처리했어. 질문하지 않은 1회 제한은 필수 답변에서 제외했어. 나머지 문항은 사람 승인 평가가 아니야. B21의 인용 없는 message 사실은 주장 47개 지표에 포함되지 않아.</div><div class="toolbar"><input id="search" type="search" aria-label="질문·답변·근거 검색" placeholder="질문, 공지 제목, 답변 내용 검색"><select id="filter" aria-label="평가 결과 필터"><option value="all">전체 32문항</option><option value="failed">답변 실패 5문항</option><option value="retrieval-failed">통합 검색 미충족 5문항</option><option value="behavior">행동 평가 8문항</option><option value="citation">인용·근거 문제 2문항</option></select><span id="count" role="status" aria-live="polite">32 / 32문항</span><button id="expand">표시된 근거 모두 펼치기</button><button id="collapse">모두 접기</button></div><div id="cases">'''+''.join(cards)+'''</div><p id="empty" class="empty" hidden>조건에 맞는 문항이 없어.</p><div class="provenance">원본: data/processed/project-benchmark-v1-resumed/report.json<br>SHA-256: '''+digest+'''<br>검토: eval/project_benchmark_review_a_v1.json / project_benchmark_review_b_v1.json<br>정정 이력: eval/project_benchmark_amendments_v1.json (B10 판정·기대 답변·기준 수정)<br>질문·답변·근거 본문은 보존된 평가 기록 그대로 표시했어. 이 보고서를 만들면서 답변을 다시 생성하지 않았어.</div></main><script>
const cards=[...document.querySelectorAll('.case')];const search=document.querySelector('#search'),filter=document.querySelector('#filter');
const haystacks=new Map(cards.map(c=>[c,c.textContent.toLocaleLowerCase()]));
function apply(){const q=search.value.trim().toLocaleLowerCase(),f=filter.value;let shown=0;for(const c of cards){const match=f==='all'||c.getAttribute('data-'+f)==='true';c.hidden=!(match&&haystacks.get(c).includes(q));if(!c.hidden)shown++;}document.querySelector('#count').textContent=shown+' / 32문항';document.querySelector('#empty').hidden=shown!==0;}
search.addEventListener('input',apply);filter.addEventListener('change',apply);
document.querySelector('#expand').onclick=()=>cards.filter(c=>!c.hidden).forEach(c=>c.querySelectorAll('.evidence').forEach(d=>d.open=true));
document.querySelector('#collapse').onclick=()=>document.querySelectorAll('details').forEach(d=>d.open=false);
document.querySelectorAll('.cite').forEach(b=>b.onclick=()=>{const d=document.getElementById(b.dataset.target);d.open=true;d.scrollIntoView({behavior:'smooth',block:'center'});});
</script></body></html>'''


def make_excel(report,digest,path):
    wb=Workbook();wb.remove(wb.active)
    summary=wb.create_sheet('평가 요약');summary.append(['항목','결과','해석'])
    for row in METRICS:summary.append(row)
    summary.append(['평가 대상','2026-10-11 원본 · 32문항','답변 가능24 / 행동8. 이후 분해검색 실험과 별도.'])
    summary.append(['읽는 법','질문별 상세 → 검색 Top5 → 주장별 검토','검색 Top5에서 질문ID로 필터하면 해당 질문의 근거5개를 볼 수 있어.'])
    summary.append(['검색 판정','Top5 전체를 합쳐 정답 전체를 지원하면 성공','B07·B10·B15·B17을 검색 성공으로 정정. 개별 묶음 일치는 참고용.'])
    summary.append(['인용 지표 주의','B21 message 사실은 claim 분모에서 제외','인용 번호가 유효하다고 모든 주장이 지원되는 것은 아니야.'])
    summary.append(['검토자','assistant 원문 대조','B10은 2026-10-12 사용자 정정으로 정답 처리. 나머지는 사람 승인 평가가 아님.'])
    summary.append(['원본 SHA-256',digest,'원본 보고서와 기존 리뷰의 해시 일치 확인'])
    questions=wb.create_sheet('질문별 상세')
    questions.append(['질문ID','유형','질문','최종 판정','응답 상태','모델 답변 원문','정답·기대 행동','판정 기준','검토 이유','완전 근거 첫 순위','실제 인용 순위','인용·근거 문제','Rank1 제목','Rank2 제목','Rank3 제목','Rank4 제목','Rank5 제목','정답 원문 경로','입력 토큰','출력 토큰','생성 시간(초)','Top5 통합 검색 판정','통합 검색 검토 이유'])
    evidence=wb.create_sheet('검색 Top5')
    evidence.append(['질문ID','질문','순위','공지 제목','파일명','모델 실제 인용','단일 묶음 정답 일치','벡터 유사도(확률 아님)','E5 토큰','검색 본문 전체','모델 전달 텍스트','공식 공지 URL','첨부 URL','원문 위치','인접 문맥 보완','청크ID','원문 구간'])
    claims=wb.create_sheet('주장별 검토')
    claims.append(['질문ID','주장 번호','주장 원문','인용 순위','근거 지원 여부','검토 이유','비고'])
    for row in report['results']:
        gen=row['generation']['result']; review=row['review']; sources=row['current']['top5']
        cited=sorted({i for c in gen['claims'] for i in c['citations']})
        grade=('완전 정답' if row['kind']=='answerable' else '행동 통과') if row['passed'] else '완전 답변 실패'
        questions.append([row['id'],KIND[row['kind']],row['query'],grade,STATUS[gen['status']],gen['answer'],row['answer_note'],row.get('rubric',''),review['reason'],row['current']['first_relevant_rank'] or ('행동 평가' if row['kind']!='answerable' else '단일 묶음 일치 없음'),', '.join(map(str,cited)) or '없음','있음' if row['citation_issue'] else '없음',*[b['document_title'] for b in sources],'; '.join(s['path'] for s in row.get('gold_sources',[])),gen['usage']['prompt_token_count'],gen['usage']['candidates_token_count'],round(row['generation']['wall_ms']/1000,3),('성공' if row['retrieval_review']['complete'] else '미충족') if row['retrieval_review'] else '행동 평가',row['retrieval_review']['reason'] if row['retrieval_review'] else ''])
        for b in sources:
            evidence.append([row['id'],row['query'],b['rank'],b['document_title'],b.get('filename') or '공지 본문','인용함' if b['rank'] in cited else '인용 안 함',('일치' if b['is_evidence'] else '미충족') if row['kind']=='answerable' else '채점 대상 아님',b['score'],b['token_count'],b['body'],b['text'],b.get('source_url'),b.get('attachment_url'),location(b),'보완함' if b['context_expanded'] else '원래 청크',' / '.join(b['chunk_ids']),json.dumps(b['source_spans'],ensure_ascii=False)])
        for i,(claim,checked) in enumerate(zip(gen['claims'],review['claim_checks']),1):
            claims.append([row['id'],i,claim['text'],', '.join(map(str,claim['citations'])),'지원됨' if checked['supported'] else '지원 안 됨',checked['reason'],'claims 배열'])
        if review.get('message_fact_leak'):
            claims.append([row['id'],'message',gen['message'],'없음','사실은 근거 있음 / 인용 누락',review['reason'],'47개 claim 지표에서 제외된 사실'])
    widths={summary:[34,40,80],questions:[12,20,60,18,23,85,65,65,90,25,18,18,40,40,40,40,40,45,14,14,18,24,90],evidence:[12,55,8,45,35,18,24,20,12,110,110,45,45,25,18,50,70],claims:[12,14,90,18,28,100,35]}
    for ws in wb:
        ws.freeze_panes='D2' if ws==questions else 'A2'
        ws.auto_filter.ref=ws.dimensions;ws.sheet_view.zoomScale=85
        for i,width in enumerate(widths[ws],1):ws.column_dimensions[get_column_letter(i)].width=width
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='17354A');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=34
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height=90 if ws!=summary else 45
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                cell.font=Font(name='맑은 고딕',size=10)
                if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F0F5F9')
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
    for row in evidence.iter_rows(min_row=2):
        for index in [11,12]:
            if href(row[index].value):row[index].hyperlink=row[index].value;row[index].font=Font(color='075FA5',underline='single')
    for row in questions.iter_rows(min_row=2):
        if row[3].value=='완전 답변 실패':row[3].fill=PatternFill('solid',fgColor='FDE8E6');row[3].font=Font(color='9F251A',bold=True)
    wb.save(path)


def main():
    report,digest=load();OUT.mkdir(exist_ok=True)
    html_path=OUT/(STEM+'.html');xlsx_path=OUT/(STEM+'.xlsx')
    if html_path.exists() or xlsx_path.exists():raise FileExistsError('Report outputs already exist')
    html_path.write_text(make_html(report,digest),encoding='utf-8')
    make_excel(report,digest,xlsx_path)
    print(html_path);print(xlsx_path)

if __name__=='__main__':main()
