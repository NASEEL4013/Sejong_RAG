"""Topic groups in DOCX forms, with verified original merged-cell geometry."""
import re
from docx import Document
from sejong_rag.extract_attachments import clean_text

class FormLayoutReview(ValueError):
    pass

def compact(text):
    return re.sub(r'\s+', '', text)

def candidate(records):
    return (all(r.get('file_type')=='docx' for r in records)
            and any('신청서' in str(r.get('filename','')) for r in records)
            and any('첨부서류' in compact(r['text']) and '확인표' in compact(r['text']) for r in records)
            and any(re.search(r'\d+자\s*이내',r['text']) for r in records))

def topic_plans(path,records):
    doc=Document(path)
    if len(doc.tables)!=1:
        raise FormLayoutReview('N_multiple_tables_require_review')
    lookup={}
    for r in records:
        loc=r['locator']
        key=('row',loc.get('table'),loc.get('row')) if 'table' in loc else ('paragraph',loc.get('paragraph'))
        if key in lookup:raise FormLayoutReview('N_duplicate_locator')
        lookup[key]=r
    common=[];mapped=set()
    for n,p in enumerate(doc.paragraphs,1):
        text=clean_text(p.text)
        if not text:continue
        r=lookup.get(('paragraph',n))
        if r is None or compact(r['text'])!=compact(text):raise FormLayoutReview('N_paragraph_mismatch')
        common.append(r);mapped.add(r['record_id'])
    if len(common)!=1:
        raise FormLayoutReview('N_common_context_requires_review')
    # Do not silently move an unrelated trailing paragraph into common context.
    from docx.oxml.ns import qn
    blocks=[node.tag for node in doc.element.body if node.tag in {qn('w:p'),qn('w:tbl')}]
    table_pos=blocks.index(qn('w:tbl'))
    if any(p.text.strip() for p in doc.paragraphs[table_pos:]):
        raise FormLayoutReview('N_trailing_context_requires_review')
    groups=[];current=[]
    for n,row in enumerate(doc.tables[0].rows,1):
        cells=list(row.cells)
        if any(c._tc.xpath('.//w:tbl') for c in cells):raise FormLayoutReview('N_nested_table')
        r=lookup.get(('row',1,n))
        if r is None or compact(r['text'])!=compact(clean_text(' | '.join(c.text for c in cells))):
            raise FormLayoutReview('N_row_source_mismatch')
        mapped.add(r['record_id']);seen=set();unique=[]
        for col,c in enumerate(cells,1):
            origin=c._tc.getroottree().getpath(c._tc)
            if origin in seen:continue
            seen.add(origin)
            unique.append(dict(column=col,columns=c._tc.grid_span,origin=origin,text=clean_text(c.text)))
        full=len(unique)==1
        value=unique[0]['text'] if full else ''
        limit=full and bool(re.fullmatch(r'[-\s]*띄어쓰기\s*(?:제외|포함)\s*\d+자\s*이내[-\s]*',value))
        if limit and not current:raise FormLayoutReview('N_limit_without_question')
        heading=full and bool(re.search(r'(?:정보|현황|경로|확인표|안내)\s*(?:\([^)]*\))?$',value)
                                  or re.match(r'^\d+\.',value))
        note=full and bool(re.match(r'^(?:※|주의|조건|.*발급기간\s*[:：])',value))
        if full and value and not (heading or note or limit):
            raise FormLayoutReview('N_full_width_row_requires_review')
        if heading and not limit and current:
            groups.append(current);current=[]
        current.append(dict(record=r,cells=unique,limit=limit))
    if current:groups.append(current)
    if mapped!={r['record_id'] for r in records}:raise FormLayoutReview('N_unmapped_records')
    plans=[]
    for group in groups:
        first=group[0]['cells']
        title=first[0]['text'] if len(first)==1 else '인적사항 및 재학정보'
        question=bool(re.match(r'^\d+\.',title))
        if question and not any(row['limit'] for row in group):
            raise FormLayoutReview('N_question_limit_missing')
        if not question and any(row['limit'] for row in group):
            raise FormLayoutReview('N_limit_owner_requires_review')
        body='문서 용도: 작성용 신청서 (빈 서식)\n'+common[0]['text']+'\n항목: '+title+'\n'
        body+='서식 내용 (왼쪽에서 오른쪽 순서):\n'
        body+='\n'.join(' | '.join(c['text'] or '[작성란]' for c in row['cells']) for row in group)
        plans.append(dict(body=body,records=[row['record'] for row in group],common=common,
                          title=title,geometry=[dict(row=row['record']['locator']['row'],cells=row['cells']) for row in group]))
    return plans
