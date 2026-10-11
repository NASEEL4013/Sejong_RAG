"""다운로드한 첨부파일을 출처 위치가 보존된 검색용 텍스트로 변환한다."""
from __future__ import annotations
import argparse, io, json, re, subprocess, tempfile, zipfile
from collections import Counter
from pathlib import Path, PurePosixPath
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
from xml.etree import ElementTree
DEFAULT_DATA_DIR = Path("data")
MAX_ZIP_MEMBERS = 200
MAX_ZIP_MEMBER_BYTES = 50 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 200 * 1024 * 1024

def clean_text(value: object) -> str:
    text = str(value or "").replace("\x00", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()

def load_attachment_metadata(data_dir: Path = DEFAULT_DATA_DIR) -> list[dict[str, object]]:
    records = {}
    for folder in ("downloaded-attachments", "imported-attachments"):
        for path in sorted((data_dir / "metadata" / folder).glob("*.json")):
            item = json.loads(path.read_text(encoding="utf-8"))
            raw_path = Path(str(item["raw_path"]))
            if not raw_path.is_absolute(): raw_path = data_dir.parent / raw_path
            item["resolved_raw_path"] = str(raw_path)
            records[str(item["document_id"])] = item
    for path in sorted((data_dir / "metadata").glob("*.json")):
        notice = json.loads(path.read_text(encoding="utf-8"))
        for attachment in notice.get("processed_attachments", []):
            document_id = str(attachment["document_id"])
            if document_id in records:
                continue
            item = {**attachment, "source_url": notice.get("source_url"), "attachment_url": attachment.get("download_url"), "filename": attachment.get("filename")}
            raw_path = Path(str(item["raw_path"]))
            if not raw_path.is_absolute():
                raw_path = data_dir.parent / raw_path
            item["resolved_raw_path"] = str(raw_path)
            records[document_id] = item
    return [records[key] for key in sorted(records)]

def _record(text: object, locator: dict[str, object], method: str):
    cleaned = clean_text(text)
    if not cleaned: return None
    return {"text": cleaned, "text_length": len(cleaned), "locator": locator, "extraction_method": method}

def extract_pdf(source):
    reader, records, warnings = PdfReader(source), [], []
    for number, page in enumerate(reader.pages, 1):
        try: text = page.extract_text(extraction_mode="layout") or page.extract_text() or ""
        except Exception as exc:
            warnings.append(f"page {number}: {type(exc).__name__}: {exc}"); continue
        item = _record(text, {"page": number}, "pypdf_layout")
        if item: records.append(item)
        else: warnings.append(f"page {number}: no embedded text (OCR may be required)")
    return records, warnings

def extract_xlsx(source):
    if isinstance(source, (str, Path)):
        with open(source, 'rb') as file: header=file.read(8)
    else:
        position=source.tell(); header=source.read(8); source.seek(position)
    if header == bytes.fromhex('D0CF11E0A1B11AE1'):
        return extract_xls(source)
    book, records = load_workbook(source, read_only=True, data_only=False), []
    for sheet in book.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            cells = [f"{index}열: {value}" for index, value in enumerate(row, 1) if value is not None]
            item = _record(" | ".join(cells), {"sheet": sheet.title, "row": row_number}, "openpyxl_row")
            if item: records.append(item)
    book.close(); return records, []

def extract_xls(source):
    """구형 XLS 셀을 직접 읽는다. 매크로나 수식을 실행하지 않는다."""
    import xlrd
    import olefile
    position = source.tell() if hasattr(source, 'tell') else None
    with olefile.OleFileIO(str(source) if isinstance(source, (str, Path)) else source) as container:
        if container.exists('EncryptedPackage'):
            raise ValueError('암호화된 엑셀입니다. 암호가 해제된 원본이 필요합니다.')
    if position is not None: source.seek(position)
    if isinstance(source, (str, Path)):
        book=xlrd.open_workbook(filename=str(source),on_demand=True)
    else:
        book=xlrd.open_workbook(file_contents=source.read(),on_demand=True)
    records=[]
    try:
        for sheet in book.sheets():
            for number in range(sheet.nrows):
                cells=[f'{column}열: {value}' for column,value in enumerate(sheet.row_values(number),1) if value!='']
                item=_record(' | '.join(cells),{'sheet':sheet.name,'row':number+1},'xlrd_row')
                if item: records.append(item)
    finally:
        book.release_resources()
    return records, []

def extract_docx(source):
    document, records = Document(source), []
    for number, paragraph in enumerate(document.paragraphs, 1):
        item = _record(paragraph.text, {"paragraph": number}, "python_docx_paragraph")
        if item: records.append(item)
    for table_number, table in enumerate(document.tables, 1):
        for row_number, row in enumerate(table.rows, 1):
            item = _record(" | ".join(cell.text for cell in row.cells), {"table": table_number, "row": row_number}, "python_docx_table")
            if item: records.append(item)
    return records, []

def extract_pptx(source):
    presentation, records = Presentation(source), []
    for slide_number, slide in enumerate(presentation.slides, 1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False): parts.append(shape.text)
            if getattr(shape, "has_table", False): parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
        item = _record("\n".join(parts), {"slide": slide_number}, "python_pptx_slide")
        if item: records.append(item)
    return records, []

def extract_hwp(path: Path):
    executable = Path(__file__).resolve().parents[2] / ".venv/bin/hwp5txt"
    command = str(executable) if executable.is_file() else "hwp5txt"
    result = subprocess.run([command, str(path)], capture_output=True, text=True, timeout=120, check=False)
    if result.returncode: raise ValueError(clean_text(result.stderr) or f"hwp5txt exit code {result.returncode}")
    item = _record(result.stdout, {"section": "document"}, "pyhwp_hwp5txt")
    return ([item] if item else []), ([] if item else ["HWP parser returned no text"])

def extract_hwpx(source):
    """HWPX의 본문 문단과 표 셀을 중복 없이 읽고 위치를 보존한다."""
    records, warnings = [], []
    def tag(node): return node.tag.rsplit("}", 1)[-1]
    def paragraph_text(node):
        if tag(node) == "t": return node.text or ""
        return "".join(paragraph_text(child) for child in node if tag(child) not in {"p", "tbl"})
    with zipfile.ZipFile(source) as archive:
        members = _safe_zip_members(archive)
        sections = sorted((m for m in members if re.fullmatch(r"Contents/section\d+\.xml", m.filename)),
                          key=lambda m: int(re.search(r"section(\d+)", m.filename).group(1)))
        for section in sections:
            payload = archive.read(section)
            if b"<!DOCTYPE" in payload.upper() or b"<!ENTITY" in payload.upper():
                raise ValueError("HWPX external/entity declarations are not supported")
            root = ElementTree.fromstring(payload)
            counts = Counter()
            def walk(node, locator):
                kind = tag(node)
                locator = dict(locator)
                if kind == "tbl":
                    counts["table"] += 1
                    locator["table"] = counts["table"]
                if kind == "p":
                    counts["paragraph"] += 1
                    locator["paragraph"] = counts["paragraph"]
                    item = _record(paragraph_text(node), locator, "hwpx_xml_paragraph")
                    if item: records.append(item)
                if kind == "pic": warnings.append(f"{section.filename}: embedded image requires visual review")
                rows, cells = 0, 0
                for child in node:
                    child_locator = dict(locator)
                    if tag(child) == "tr":
                        rows += 1; child_locator["row"] = rows
                    if tag(child) == "tc":
                        cells += 1; child_locator["column"] = cells
                    walk(child, child_locator)
            walk(root, {"section": section.filename})
    if not sections: warnings.append("HWPX has no body sections")
    return records, list(dict.fromkeys(warnings))

def _safe_zip_members(archive):
    members = [item for item in archive.infolist() if not item.is_dir()]
    if len(members) > MAX_ZIP_MEMBERS: raise ValueError(f"ZIP member limit exceeded: {len(members)}")
    total = 0
    for item in members:
        path = PurePosixPath(item.filename)
        if path.is_absolute() or ".." in path.parts: raise ValueError(f"unsafe ZIP path: {item.filename}")
        if item.file_size > MAX_ZIP_MEMBER_BYTES: raise ValueError(f"ZIP member size limit exceeded: {item.filename}")
        total += item.file_size
    if total > MAX_ZIP_TOTAL_BYTES: raise ValueError(f"ZIP expanded size limit exceeded: {total}")
    return members

def _zip_display_name(name):
    try:
        return name.encode("cp437").decode("cp949")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return name

def extract_zip(source, depth=0):
    if depth > 3: return [], ["nested ZIP depth limit exceeded"]
    records, warnings = [], []
    with zipfile.ZipFile(source) as archive:
        for member in _safe_zip_members(archive):
            name = _zip_display_name(member.filename)
            suffix = Path(name).suffix.lower().lstrip(".")
            payload = archive.read(member)
            try:
                if suffix == "txt":
                    item = _record(payload.decode("utf-8-sig", errors="replace"), {}, "zip_text")
                    nested, nested_warnings = ([item] if item else []), []
                elif suffix == "hwp":
                    with tempfile.NamedTemporaryFile(suffix=".hwp") as temporary:
                        temporary.write(payload); temporary.flush()
                        nested, nested_warnings = extract_hwp(Path(temporary.name))
                elif suffix == "zip":
                    nested, nested_warnings = extract_zip(io.BytesIO(payload), depth + 1)
                elif suffix in EXTRACTORS:
                    nested, nested_warnings = EXTRACTORS[suffix](io.BytesIO(payload))
                else:
                    warnings.append(f"{name}: unsupported ZIP member type"); continue
                for item in nested:
                    item["locator"] = {"zip_member": name, **item["locator"]}; records.append(item)
                warnings.extend(f"{name}: {warning}" for warning in nested_warnings)
            except Exception as exc:
                warnings.append(f"{name}: {type(exc).__name__}: {exc}")
    return records, warnings
EXTRACTORS = {"pdf": extract_pdf, "xlsx": extract_xlsx, "xlsm": extract_xlsx, "xls": extract_xls,
              "hwpx": extract_hwpx, "docx": extract_docx, "pptx": extract_pptx}

def extract_one(metadata, output_dir: Path, refresh=False):
    document_id, raw_path = str(metadata["document_id"]), Path(str(metadata["resolved_raw_path"]))
    file_type = str(metadata.get("file_type") or raw_path.suffix.lstrip(".")).lower()
    records_path, report_path = output_dir / "records" / f"{document_id}.jsonl", output_dir / "reports" / f"{document_id}.json"
    if report_path.is_file() and not refresh: return json.loads(report_path.read_text(encoding="utf-8")) | {"run_status": "skipped"}
    records, warnings, error = [], [], None
    try:
        if not raw_path.is_file(): raise FileNotFoundError(raw_path)
        if file_type == "hwp": records, warnings = extract_hwp(raw_path)
        elif file_type == "zip": records, warnings = extract_zip(raw_path)
        elif file_type == "doc": raise ValueError("legacy DOC extraction is not supported safely; conversion/review required")
        elif file_type in EXTRACTORS: records, warnings = EXTRACTORS[file_type](raw_path)
        else: raise ValueError(f"unsupported file type: {file_type}")
    except Exception as exc: error = f"{type(exc).__name__}: {exc}"
    status = "failed" if error else ("review_required" if warnings or not records else "success")
    records_path.parent.mkdir(parents=True, exist_ok=True)
    with records_path.open("w", encoding="utf-8") as output:
        for index, item in enumerate(records, 1):
            output.write(json.dumps({"record_id": f"{document_id}-record-{index}", "document_id": document_id, "parent_document_id": metadata.get("parent_document_id"), "filename": metadata.get("filename"), "file_type": file_type, "source_url": metadata.get("source_url"), "attachment_url": metadata.get("attachment_url"), **item}, ensure_ascii=False) + "\n")
    report = {"document_id": document_id, "filename": metadata.get("filename"), "file_type": file_type, "raw_path": str(raw_path), "records_path": str(records_path), "status": status, "record_count": len(records), "text_length": sum(int(item["text_length"]) for item in records), "warnings": warnings, "error": error}
    report_path.parent.mkdir(parents=True, exist_ok=True); report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report | {"run_status": "processed"}

def extract_all(data_dir=DEFAULT_DATA_DIR, output_dir=None, refresh=False, max_items=None):
    output_dir = output_dir or data_dir / "processed/attachments"
    metadata = load_attachment_metadata(data_dir)
    if max_items is not None: metadata = metadata[:max_items]
    reports = [extract_one(item, output_dir, refresh) for item in metadata]
    return {"attachment_count": len(reports), "status_counts": dict(Counter(str(item["status"]) for item in reports)), "format_counts": dict(Counter(str(item["file_type"]) for item in reports)), "record_count": sum(int(item["record_count"]) for item in reports), "text_length": sum(int(item["text_length"]) for item in reports), "processed": sum(item["run_status"] == "processed" for item in reports), "skipped": sum(item["run_status"] == "skipped" for item in reports)}

def main():
    parser = argparse.ArgumentParser(description="공식 첨부파일에서 출처 위치가 보존된 텍스트를 추출합니다.")
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR); parser.add_argument("--output-dir", type=Path); parser.add_argument("--max-items", type=int); parser.add_argument("--refresh", action="store_true")
    args = parser.parse_args(); print(json.dumps(extract_all(args.data_dir, args.output_dir, args.refresh, args.max_items), ensure_ascii=False, indent=2))
if __name__ == "__main__": main()
